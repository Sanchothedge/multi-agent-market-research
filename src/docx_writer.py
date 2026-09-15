"""Renders the LLM-generated Markdown report as a real .docx file.

The original workflow ("Prepare Word Report" + "Create Word-Compatible
File") built an RTF file and gave it a .rtf extension while calling it a
Word-compatible file. This writes an actual .docx via python-docx, including
real Word tables for the Markdown pipe tables the report prompt requires.
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_runs_with_bold(paragraph, text: str) -> None:
    """Adds text to a paragraph, turning **bold** spans into bold runs."""
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _is_table_separator(line: str) -> bool:
    stripped = line.strip().strip("|")
    if not stripped:
        return False
    cells = [c.strip() for c in stripped.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def _parse_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [c.strip() for c in stripped.split("|")]


def markdown_to_docx(title: str, subtitle: str, markdown: str, output_path: str) -> str:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if subtitle:
        sub = doc.add_paragraph(subtitle)
        sub.runs[0].italic = True
        sub.runs[0].font.size = Pt(10)
    doc.add_paragraph()

    lines = markdown.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Markdown table: a header row followed by a separator row.
        if stripped.startswith("|") and i + 1 < n and _is_table_separator(lines[i + 1]):
            header_cells = _parse_table_row(stripped)
            i += 2
            body_rows: list[list[str]] = []
            while i < n and lines[i].strip().startswith("|"):
                body_rows.append(_parse_table_row(lines[i]))
                i += 1

            table = doc.add_table(rows=1, cols=len(header_cells))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for idx, cell_text in enumerate(header_cells):
                hdr_cells[idx].text = cell_text
                for p in hdr_cells[idx].paragraphs:
                    for r in p.runs:
                        r.bold = True
            for row_values in body_rows:
                cells = table.add_row().cells
                for idx, cell_text in enumerate(row_values):
                    if idx < len(cells):
                        cells[idx].text = cell_text
            doc.add_paragraph()
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, re.sub(r"^\d+\.\s+", "", stripped))
        else:
            p = doc.add_paragraph()
            _add_runs_with_bold(p, stripped)

        i += 1

    doc.save(output_path)
    return output_path
